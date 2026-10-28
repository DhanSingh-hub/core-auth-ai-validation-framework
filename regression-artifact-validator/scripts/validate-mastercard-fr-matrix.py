"""Validate the Mastercard Credit Purchase FR matrix workbook against the current Test Solution.

The original workbook is hash-bound and parsed as BR -> FR -> TS -> validation lines -> expected results. Every
line must have exactly one curated Test Solution assessment (docs/specs/kb/segment-100/
mastercard-fr-matrix-validation.json). The script then checks those assessments against:

* ATL105 Chapter 13 element names (every "Element N" citation in the workbook),
* resolvable ATL105 source anchors and section-absence claims,
* the canonical Test Solution package and the Mastercard Credit Purchase training package (live BR status),
* the Section 10.1 credit-card draft supplement and the Mastercard BA reference review (lessons, questions,
  dispositions and screenshot-transcription fidelity).

Outcomes follow the handbook's Independent Review Process. Nothing is approved, promoted or executed.

Usage, from the validator module:  python scripts/validate-mastercard-fr-matrix.py [--check]
"""
import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
MODULE = HERE.parent
PACK = MODULE / "specifications" / "ATL105"
KB = PACK / "docs" / "specs" / "kb" / "segment-100"
CURATED = KB / "mastercard-fr-matrix-validation.json"
BA_REVIEW = KB / "mastercard-ba-reference-review.json"
REVIEW_DIR = PACK / "test-output" / "test-solution-independent-review"
TRAINING = PACK / "test-output" / "test-json" / "mc-credit-purchase-training-package.json"
DRAFT_10_1 = REVIEW_DIR / "section-10-1-credit-card-processing-review-v4" / "section-10-1-br-ts-tc-td-draft-package.json"
SPEC = PACK / "docs" / "specs" / "extracted_text.txt"
OUT_DIR = REVIEW_DIR / "mc-fr-matrix-validation-v1"
TEXT_OUTPUTS = ("README.md", "validation.json", "line-validation.csv", "findings.csv")

HEADER = ["BR ID", "Business Requirement", "FR ID", "Functional Requirement", "TS ID", "Test Scenario",
          "Acceptance Criteria", "Validations", "Expected Result"]
VERDICTS = ("SUPPORTED", "SUPPORTED_WITH_CORRECTION", "CONTRADICTED", "EXTERNAL_POLICY", "NOT_TESTABLE_AS_WRITTEN")
SUPPORTED_VERDICTS = ("SUPPORTED", "SUPPORTED_WITH_CORRECTION")
REVIEW_VERDICTS = ("EXTERNAL_POLICY", "NOT_TESTABLE_AS_WRITTEN")
CLASSES = ("SOURCE_CONFLICT", "TEST_SOLUTION_GAP", "AI_ARTIFACT_DEFECT", "CHAIN_GAP", "PROVENANCE_OR_VERSION_GAP",
           "REVIEW_REQUIRED")
SEVERITIES = ("HIGH", "MEDIUM", "LOW", "NONE")
EXPECTED_VERDICTS = ("MESSAGE_LEVEL", "SOURCE_SUPPORTED_NEEDS_EXECUTION", "NEEDS_EXECUTION_EVIDENCE", "EXTERNAL_CONTRACT")
OUTCOMES = ("REJECTED_NOT_INTAKE_READY", "REVIEW_REQUIRED", "ACCEPTED_WITH_GAPS", "ACCEPTED")
ARTIFACT_OWNER = "BA (SEG100-SME-003 questionnaire)"
TEST_SOLUTION_OWNER = "Test Solution team"


def long_path(path):
    """Windows extended-length form; the Section 10.1 package path exceeds 260 characters (handbook L25)."""
    resolved = os.path.abspath(str(path))
    if os.name == "nt" and not resolved.startswith("\\\\?\\"):
        return "\\\\?\\" + resolved
    return resolved


def read_bytes(path):
    with open(long_path(path), "rb") as handle:
        return handle.read()


def read_json(path):
    return json.loads(read_bytes(path).decode("utf-8-sig"))


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def rel(path):
    resolved = Path(os.path.abspath(str(path)))
    try:
        return resolved.relative_to(MODULE).as_posix()
    except ValueError:
        return "../" + resolved.relative_to(MODULE.parent).as_posix()


def load_module(name, filename):
    spec = importlib.util.spec_from_file_location(name, HERE / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def worst(values, order):
    return min(values, key=order.index) if values else order[-1]


# ---------------------------------------------------------------- workbook

def parse_workbook(path, sheet):
    """Parse the FR matrix; rows are BR -> FR -> TS -> validation lines, as in the approved template."""
    workbook = openpyxl.load_workbook(long_path(path), data_only=True)
    if sheet not in workbook.sheetnames:
        raise ValueError(f"Workbook has no sheet {sheet!r}: {workbook.sheetnames}")
    ws = workbook[sheet]
    header = [ws.cell(2, column).value for column in range(1, 10)]
    if header != HEADER:
        raise ValueError(f"Unexpected FR matrix header: {header}")
    brs, frs, scenarios, structure, owners = [], [], [], [], {}
    fr = scenario = None
    for r in range(3, ws.max_row + 1):
        row = [ws.cell(r, column).value for column in range(1, 10)]
        if all(value in (None, "") for value in row):
            continue
        for column, value in enumerate(row, start=1):
            if isinstance(value, str) and value != value.strip():
                structure.append({"check": "LEADING_OR_TRAILING_WHITESPACE", "cell": f"{get_column_letter(column)}{r}",
                                  "text": value})
        if row[0]:
            brs.append({"id": str(row[0]).strip(), "text": str(row[1] or "").strip(), "cell": f"A{r}"})
            owners.update({f"A{r}": brs[-1]["id"], f"B{r}": brs[-1]["id"]})
        if row[2]:
            fr = {"id": str(row[2]).strip(), "name": str(row[3] or "").strip(), "scenarios": []}
            frs.append(fr)
            owners.update({f"C{r}": fr["id"], f"D{r}": fr["id"]})
        if row[4]:
            if fr is None:
                raise ValueError(f"Scenario before any FR in row {r}")
            scenario = {"id": str(row[4]).strip(), "fr": fr["id"], "name": str(row[5] or "").strip(),
                        "ac": str(row[6] or "").strip(), "expected": str(row[8] or "").strip(), "row": r, "lines": []}
            scenarios.append(scenario)
            fr["scenarios"].append(scenario["id"])
            owners.update({f"{column}{r}": scenario["id"] for column in "EFGI"})
        elif any(row[i] for i in (5, 6, 8)):
            raise ValueError(f"Scenario content without a TS ID in row {r}")
        if row[7] is not None:
            if scenario is None:
                raise ValueError(f"Validation before any scenario in row {r}")
            text = str(row[7])
            line = {"id": f"{scenario['id']}.V{len(scenario['lines']) + 1}", "text": text, "cell": f"H{r}"}
            scenario["lines"].append(line)
            owners[line["cell"]] = line["id"]
            if re.search(r";|,\s", text):
                structure.append({"check": "MULTIPLE_CHECKS_IN_ONE_CELL", "cell": line["cell"], "line": line["id"],
                                  "text": text})
    if len(brs) != 1:
        raise ValueError(f"Expected exactly one business requirement, found {len(brs)}")
    for item in structure:
        item.setdefault("line", owners.get(item["cell"], item["cell"]))
    for item in scenarios:
        for field, check in (("name", "MISSING_SCENARIO_NAME"), ("ac", "MISSING_ACCEPTANCE_CRITERIA"),
                             ("expected", "MISSING_EXPECTED_RESULT")):
            if not item[field]:
                structure.append({"check": check, "cell": f"row {item['row']}", "line": item["id"], "text": ""})
        if not item["lines"]:
            structure.append({"check": "NO_VALIDATION_LINE", "cell": f"row {item['row']}", "line": item["id"], "text": ""})
    ids = [item["id"] for item in scenarios]
    if ids != [f"TS-{n:03d}" for n in range(1, len(ids) + 1)] or len(set(ids)) != len(ids):
        structure.append({"check": "TS_ID_SEQUENCE", "cell": "column E", "text": ", ".join(ids)})
    structure.append({"check": "NO_VALIDATION_LINE_IDS", "cell": "column H",
                      "text": "Validation lines have no IDs; this review assigns <TS ID>.V<n> in sheet order."})
    return {"br": brs[0], "frs": frs, "scenarios": scenarios, "structure": structure, "maxRow": ws.max_row}


# ---------------------------------------------------------------- ATL105 source

def chapter13_names(lines):
    names = {}
    for number, text in enumerate(lines, 1):
        match = re.match(r"^\s*Number:\s+(\d+)\s+Name:\s+(.+?)\s*$", text)
        if match and int(match.group(1)) not in names:
            names[int(match.group(1))] = {"line": number, "name": match.group(2)}
    return names


def absence_evidence(lines, spec):
    starts = [i for i, text in enumerate(lines) if re.match(rf"^\s*{re.escape(spec['section'])}", text)]
    if not starts:
        raise ValueError(f"Section heading not found: {spec['section']}")
    start = starts[-1]
    end = next((i for i in range(start + 1, len(lines)) if re.match(rf"^\s*{re.escape(spec['next'])}\s+\S", lines[i])), None)
    if end is None:
        raise ValueError(f"Section end not found after {spec['section']}")
    found = [start + 1 + i for i, text in enumerate(lines[start:end]) if spec["term"].lower() in text.lower()]
    if found:
        raise ValueError(f"Absence claim no longer holds: '{spec['term']}' occurs in {spec['section']} at {found}")
    return {"section": spec["section"], "lines": f"{start + 1}-{end}", "term": spec["term"], "occurrences": 0}


# ---------------------------------------------------------------- evaluation

def check_curated_line(line):
    problems = []
    if line["verdict"] not in VERDICTS:
        problems.append("unknown verdict")
    if set(line["classes"]) - set(CLASSES) or line["severity"] not in SEVERITIES:
        problems.append("unknown class or severity")
    if not line["assessment"].strip():
        problems.append("missing assessment")
    if not (line["sources"] or line.get("absence") or line["brIds"] or line["draftBrIds"]):
        problems.append("no evidence")
    if line["verdict"] == "SUPPORTED" and (line["classes"] or line["severity"] != "NONE"):
        problems.append("SUPPORTED line carries a finding")
    if line["verdict"] in ("CONTRADICTED", "SUPPORTED_WITH_CORRECTION"):
        if "AI_ARTIFACT_DEFECT" not in line["classes"] or not line["correction"].strip():
            problems.append("defect line needs AI_ARTIFACT_DEFECT and a correction")
    if line["verdict"] == "CONTRADICTED" and line["severity"] != "HIGH":
        problems.append("CONTRADICTED line must be HIGH")
    if line["verdict"] in REVIEW_VERDICTS and ("REVIEW_REQUIRED" not in line["classes"] or not line["questions"]):
        problems.append("review line needs REVIEW_REQUIRED and an open question")
    if problems:
        raise ValueError(f"Curated assessment {line['id']} is invalid: {'; '.join(problems)}")


def line_coverage(states, drafts):
    if "TRAINED" in states:
        return "TRAINED_CHAIN"
    if "REVIEW_QUEUE" in states:
        return "REVIEW_QUEUE_ONLY"
    if states:
        return "OUTSIDE_MC_CONTEXT"
    return "DRAFT_ONLY" if drafts else "NONE"


def scenario_outcome(lines, expected):
    verdicts = {line["verdict"] for line in lines}
    classes = {name for line in lines for name in line["classes"]}
    if "CONTRADICTED" in verdicts:
        return "REJECTED_NOT_INTAKE_READY"
    if verdicts & set(REVIEW_VERDICTS) or "REVIEW_REQUIRED" in classes:
        return "REVIEW_REQUIRED"
    if (verdicts != {"SUPPORTED"} or any(line["coverage"] != "TRAINED_CHAIN" for line in lines)
            or expected["verdict"] != "MESSAGE_LEVEL"):
        return "ACCEPTED_WITH_GAPS"
    return "ACCEPTED"


def build_report(workbook_path=None):
    curated = read_json(CURATED)
    ba = read_json(BA_REVIEW)
    workbook_path = Path(workbook_path) if workbook_path else (MODULE / curated["artifact"]["path"]).resolve()
    workbook_bytes = read_bytes(workbook_path)
    if sha256(workbook_bytes) != curated["artifact"]["sha256"]:
        raise ValueError("Workbook content changed since the curated review (SHA-256 mismatch); re-review it before use.")
    parsed = parse_workbook(workbook_path, curated["artifact"]["sheet"])

    workbook_lines = [(line["id"], line["text"]) for item in parsed["scenarios"] for line in item["lines"]]
    curated_lines = [(line["id"], line["text"]) for line in curated["lines"]]
    if workbook_lines != curated_lines:
        drift = [(w, c) for w, c in zip(workbook_lines, curated_lines) if w != c][:5]
        raise ValueError(f"Workbook lines differ from the curated assessment ({len(workbook_lines)} vs "
                         f"{len(curated_lines)}): {drift}")
    workbook_expected = [(item["id"], item["expected"]) for item in parsed["scenarios"]]
    if workbook_expected != [(item["ts"], item["text"]) for item in curated["expectedResults"]]:
        raise ValueError("Workbook expected results differ from the curated assessment")
    for item in curated["expectedResults"]:
        if item["verdict"] not in EXPECTED_VERDICTS or not item["note"].strip():
            raise ValueError(f"Invalid expected-result assessment for {item['ts']}")

    spec_lines = read_bytes(SPEC).decode("utf-8-sig").replace("\r\n", "\n").splitlines()
    names = chapter13_names(spec_lines)
    ba_module = load_module("mc_ba_review", "build-mastercard-reference-review.py")
    if set(curated["sources"]) & set(ba["sources"]):
        raise ValueError("Curated source IDs collide with BA review source IDs")
    sources = ba_module.resolve_sources({**ba["sources"], **curated["sources"]}, PACK)

    mc = load_module("mc_fr_matrix", "build-mastercard-fr-matrix.py")
    ts_fr = mc.ts_fr
    _, brs, _, _, _, br_text, _, _, _ = ts_fr.load()
    complete_ids = {b["id"] for b in brs}
    training = read_json(TRAINING)
    trained = {b["id"] for b in training["businessRequirements"]}
    review_queue = {item["id"]: item for item in training["reviewQueue"]}
    template_status = {item["id"]: item for item in training["businessContext"]["scenarios"]}
    drafts = {b["id"]: b for b in read_json(DRAFT_10_1)["businessRequirements"]}
    lessons = {item["id"]: item for item in ba["lessons"]}
    questions = {item["id"]: item for item in ba["questions"]}
    ba_scenarios = {item["id"]: item for item in ba["scenarios"]}

    def br_state(br_id):
        if br_id in trained:
            return "TRAINED"
        if br_id in review_queue:
            return "REVIEW_QUEUE"
        if br_id in complete_ids:
            return "NOT_IN_MC_CONTEXT"
        raise ValueError(f"Unknown Test Solution BR cited: {br_id}")

    lines_by_id = {}
    for line in curated["lines"]:
        check_curated_line(line)
        for key, index in (("lessons", lessons), ("questions", questions), ("draftBrIds", drafts), ("sources", sources)):
            unknown = set(line[key]) - set(index)
            if unknown:
                raise ValueError(f"{line['id']} cites unknown {key}: {sorted(unknown)}")
        states = OrderedDict((b, br_state(b)) for b in line["brIds"])
        related = OrderedDict((b, br_state(b)) for b in line["relatedBrIds"])
        coverage = line_coverage(set(states.values()), line["draftBrIds"])
        classes = list(line["classes"])
        if line["verdict"] in SUPPORTED_VERDICTS:
            if coverage in ("OUTSIDE_MC_CONTEXT", "DRAFT_ONLY", "NONE") and "TEST_SOLUTION_GAP" not in classes:
                classes.append("TEST_SOLUTION_GAP")
            if coverage == "REVIEW_QUEUE_ONLY" and "CHAIN_GAP" not in classes:
                classes.append("CHAIN_GAP")
        evidence = [{"sourceId": key, "section": sources[key]["section"], "path": sources[key]["path"],
                     "line": sources[key]["line"], "sha256": sources[key]["sha256"]} for key in line["sources"]]
        if line.get("absence"):
            evidence.append({"absence": absence_evidence(spec_lines, line["absence"]), "path": "docs/specs/extracted_text.txt"})
        lines_by_id[line["id"]] = {
            "id": line["id"], "ts": line["id"].split(".")[0], "text": line["text"], "verdict": line["verdict"],
            "severity": line["severity"], "classes": classes, "curatedClasses": list(line["classes"]),
            "assessment": line["assessment"], "correction": line["correction"], "evidence": evidence,
            "testSolutionBrs": [{"id": k, "state": v} for k, v in states.items()],
            "relatedBrs": [{"id": k, "state": v} for k, v in related.items()],
            "draftBrs": [{"id": d, "section": drafts[d]["sourceAnchors"][0].get("section"), "title": drafts[d]["title"]}
                         for d in line["draftBrIds"]],
            "coverage": coverage, "lessons": line["lessons"], "questions": line["questions"]}

    element_check = []
    claims = {(item["line"], item["cited"]): item for item in curated["elementClaims"]}
    mentions = set()
    for item in parsed["scenarios"]:
        for line in item["lines"]:
            for number in re.findall(r"\bElement\s+(\d+)\b", line["text"]):
                mentions.add((line["id"], int(number)))
        for field in ("ac", "expected", "name"):
            if re.search(r"\bElement\s+\d+\b", item[field]):
                raise ValueError(f"{item['id']} cites an element in its {field}; add an element claim for it")
    if mentions != set(claims):
        raise ValueError(f"Element citations and claims differ: workbook {sorted(mentions)} vs claims {sorted(claims)}")
    for (line_id, cited), claim in sorted(claims.items()):
        cited_def, correct_def = names.get(cited), names.get(claim["correct"])
        if not cited_def or not correct_def:
            raise ValueError(f"Chapter 13 definition missing for {line_id}")
        if claim["correctNameContains"] not in correct_def["name"]:
            raise ValueError(f"Element {claim['correct']} is no longer '{claim['correctNameContains']}'")
        result = "MATCH" if cited == claim["correct"] else "MISMATCH"
        if (result == "MISMATCH") != (lines_by_id[line_id]["verdict"] == "CONTRADICTED"):
            raise ValueError(f"{line_id}: element check {result} disagrees with verdict {lines_by_id[line_id]['verdict']}")
        element_check.append({"line": line_id, "text": lines_by_id[line_id]["text"], "subject": claim["subject"],
                              "citedElement": cited, "citedName": cited_def["name"], "citedSpecLine": cited_def["line"],
                              "correctElement": claim["correct"], "correctName": correct_def["name"],
                              "correctSpecLine": correct_def["line"], "result": result})

    fidelity = []
    for item in parsed["scenarios"]:
        claim = ba_scenarios.get(item["id"])
        if claim is None:
            fidelity.append({"ts": item["id"], "check": "MISSING_FROM_BA_REVIEW"})
            continue
        if claim["name"] != item["name"]:
            fidelity.append({"ts": item["id"], "check": "NAME_DIFFERS", "workbook": item["name"], "baReview": claim["name"]})
        if claim["fr"] != item["fr"]:
            fidelity.append({"ts": item["id"], "check": "FR_DIFFERS", "workbook": item["fr"], "baReview": claim["fr"]})
        cited = sorted({n for line in item["lines"] for n in re.findall(r"\bElement\s+(\d+)\b", line["text"])})
        transcribed = sorted(set(re.findall(r"\bElement\s+(\d+)\b", claim["referenceClaim"])))
        if cited != transcribed:
            fidelity.append({"ts": item["id"], "check": "ELEMENTS_DIFFER", "workbook": cited, "baReview": transcribed})
    transcription = []
    for diff in curated["transcriptionDifferences"]:
        scenario = next(s for s in parsed["scenarios"] if s["id"] == diff["ts"])
        claim_text = ba_scenarios[diff["ts"]]["referenceClaim"]
        if diff["workbookText"] not in [line["text"].strip() for line in scenario["lines"]]:
            raise ValueError(f"Transcription note for {diff['ts']} cites text not in the workbook")
        if diff["baClaimExcerpt"] is not None and diff["baClaimExcerpt"] not in claim_text:
            raise ValueError(f"Transcription note for {diff['ts']} cites text not in the BA claim")
        if diff["baClaimExcerpt"] is None and diff["workbookText"].lower() in claim_text.lower():
            raise ValueError(f"Transcription note for {diff['ts']} says the line is missing, but it is transcribed")
        transcription.append({**diff, "baReferenceClaim": claim_text, "verified": True})

    expected = {item["ts"]: item for item in curated["expectedResults"]}
    scenarios = []
    for item in parsed["scenarios"]:
        rows = [lines_by_id[line["id"]] for line in item["lines"]]
        ba_row = ba_scenarios.get(item["id"], {})
        status = template_status.get(item["id"], {})
        scenarios.append({
            "id": item["id"], "fr": item["fr"], "name": item["name"], "acceptanceCriteria": item["ac"],
            "expectedResult": item["expected"], "expectedVerdict": expected[item["id"]]["verdict"],
            "expectedNote": expected[item["id"]]["note"], "lineIds": [row["id"] for row in rows],
            "outcome": scenario_outcome(rows, expected[item["id"]]),
            "worstLineVerdict": worst([row["verdict"] for row in rows], ("CONTRADICTED", "NOT_TESTABLE_AS_WRITTEN",
                                                                         "EXTERNAL_POLICY", "SUPPORTED_WITH_CORRECTION",
                                                                         "SUPPORTED")),
            "baDisposition": ba_row.get("disposition"), "baCorrectedRequirement": ba_row.get("correctedRequirement"),
            "templateCoverage": status.get("coverageStatus"), "trainingState": status.get("trainingState"),
            "mappedTestSolutionBrIds": list(mc.TEMPLATE_MAP.get(item["id"], ([],))[0]),
            "questions": sorted({q for row in rows for q in row["questions"]})})

    current_scope = trained | set(review_queue)
    multi_segment = []
    for b in brs:
        if b["id"] in current_scope:
            continue
        raw = {str(a.get("segment") or "") for a in ts_fr.anchors(b)}
        if not any("," in s for s in raw):
            continue
        split = {part.strip() for s in raw for part in s.split(",") if part.strip()}
        text = br_text[b["id"]][0]
        if split & mc.MC_SEGMENTS and not mc.EXCLUDE_ID.search(b["id"]) and not mc.EXCLUDE_TEXT.search(text):
            multi_segment.append({"id": b["id"], "segments": sorted(raw), "title": text})

    findings = build_findings(curated, ba, parsed, lines_by_id, scenarios, element_check, transcription, fidelity,
                              review_queue, multi_segment, mc)
    lines = [lines_by_id[line_id] for line_id, _ in workbook_lines]
    inputs = [{"path": rel(path), "sha256": sha256(read_bytes(path))}
              for path in (CURATED, BA_REVIEW, TRAINING, ts_fr.PACKAGE, DRAFT_10_1, SPEC)]
    summary = OrderedDict([
        ("businessRequirements", 1), ("functionalRequirements", len(parsed["frs"])), ("scenarios", len(scenarios)),
        ("validationLines", len(lines)), ("expectedResults", len(expected)),
        ("lineVerdicts", dict(sorted(Counter(line["verdict"] for line in lines).items()))),
        ("lineSeverities", dict(sorted(Counter(line["severity"] for line in lines).items()))),
        ("lineCoverage", dict(sorted(Counter(line["coverage"] for line in lines).items()))),
        ("findingClassesOnLines", dict(sorted(Counter(c for line in lines for c in line["classes"]).items()))),
        ("scenarioOutcomes", dict(sorted(Counter(s["outcome"] for s in scenarios).items()))),
        ("expectedResultVerdicts", dict(sorted(Counter(s["expectedVerdict"] for s in scenarios).items()))),
        ("elementCitations", len(element_check)),
        ("elementMismatches", sum(item["result"] == "MISMATCH" for item in element_check)),
        ("structureObservations", len(parsed["structure"])),
        ("transcriptionDifferences", len(transcription)), ("baFidelityMismatches", len(fidelity)),
        ("findings", len(findings)),
        ("findingsBySeverity", dict(sorted(Counter(f["severity"] for f in findings).items()))),
    ])
    return OrderedDict([
        ("artifact", "ATL105-MASTERCARD-FR-MATRIX-WORKBOOK-VALIDATION"), ("status", curated["status"]),
        ("specification", curated["specification"]), ("specificationVersion", curated["specificationVersion"]),
        ("reviewer", curated["reviewer"]), ("reviewedOn", curated["reviewedOn"]),
        ("workbook", OrderedDict([("path", rel(workbook_path)), ("sha256", sha256(workbook_bytes)),
                                  ("sheet", curated["artifact"]["sheet"]), ("producer", curated["artifact"]["producer"]),
                                  ("firstCommit", curated["artifact"]["firstCommit"]),
                                  ("authority", curated["artifact"]["authority"])])),
        ("businessRequirement", parsed["br"]), ("functionalRequirements", parsed["frs"]),
        ("outcome", worst([s["outcome"] for s in scenarios], OUTCOMES)),
        ("executionCertified", False), ("smeApproval", "NOT_ESTABLISHED_BY_THIS_REVIEW"),
        ("summary", summary), ("inputs", inputs), ("scenarios", scenarios), ("lines", lines),
        ("elementCheck", element_check), ("transcriptionDifferences", transcription), ("baFidelityChecks", fidelity),
        ("structure", parsed["structure"]), ("multiSegmentScopeExclusions", multi_segment), ("findings", findings),
    ])


def build_findings(curated, ba, parsed, lines, scenarios, element_check, transcription, fidelity, review_queue,
                   multi_segment, mc):
    all_lines = list(lines.values())

    def finding(fid, cls, severity, title, detail, affected, owner, disposition, status="OPEN"):
        return OrderedDict([("id", fid), ("class", cls), ("severity", severity), ("title", title), ("detail", detail),
                            ("affected", affected), ("owner", owner), ("disposition", disposition), ("status", status)])

    contradicted = [line for line in all_lines if line["verdict"] == "CONTRADICTED"]
    corrected = [line for line in all_lines if line["verdict"] == "SUPPORTED_WITH_CORRECTION"]
    review = [line for line in all_lines if line["verdict"] in REVIEW_VERDICTS]
    gaps = [line for line in all_lines if "TEST_SOLUTION_GAP" in line["classes"]]
    chain_lines = [line for line in all_lines if any(b["state"] == "REVIEW_QUEUE" for b in line["testSolutionBrs"])]
    queued = sorted({b["id"] for line in chain_lines for b in line["testSolutionBrs"] if b["state"] == "REVIEW_QUEUE"})
    not_message = [s for s in scenarios if s["expectedVerdict"] != "MESSAGE_LEVEL"]
    curated_findings = {item["id"]: item for item in curated["findings"]}
    mismatches = [item for item in element_check if item["result"] == "MISMATCH"]
    out = []
    a01 = curated_findings["MCV-A01"]
    out.append(finding("MCV-A01", a01["class"], a01["severity"], a01["title"], a01["detail"], a01["affected"],
                       a01["owner"], a01["disposition"]))
    out.append(finding(
        "MCV-A02", "AI_ARTIFACT_DEFECT", "HIGH", "Field mappings that contradict ATL105",
        f"{len(contradicted)} lines would test the wrong field: "
        +         "; ".join(f"{line['id']} '{line['text'].strip()}' -> {line['correction'].rstrip('.')}" for line in contradicted)
        + f". {len(mismatches)} of {len(element_check)} element citations name the wrong Chapter 13 element.",
        [line["id"] for line in contradicted], ARTIFACT_OWNER, "Correct the field mapping before any test is derived."))
    high = [line for line in corrected if line["severity"] == "HIGH"]
    out.append(finding(
        "MCV-A03", "AI_ARTIFACT_DEFECT", worst([line["severity"] for line in corrected], SEVERITIES),
        "Lines whose intent is supported but whose wording, scope or field must be corrected",
        f"{len(corrected)} lines need the correction recorded in line-validation.csv (over-general rules, labels "
        "instead of encoded values, non-atomic checks, missing context)."
        + "".join(f" HIGH: {line['id']} '{line['text'].strip()}' - {line['correction']}" for line in high),
        [line["id"] for line in corrected], ARTIFACT_OWNER, "Apply the corrected requirement for each line."))
    out.append(finding(
        "MCV-A04", "REVIEW_REQUIRED", "MEDIUM", "Checks that depend on contracts outside ATL105 or are not testable as written",
        f"{sum(l['verdict'] == 'EXTERNAL_POLICY' for l in review)} lines need issuer, BIN, merchant, routing, clearing, "
        f"settlement or audit contracts; {sum(l['verdict'] == 'NOT_TESTABLE_AS_WRITTEN' for l in review)} name no "
        "testable field or rule. Open questions: " + ", ".join(sorted({q for l in review for q in l["questions"]})) + ".",
        [line["id"] for line in review], ARTIFACT_OWNER, "Supply the contract or rule named in each open question."))
    out.append(finding(
        "MCV-A05", "REVIEW_REQUIRED", "MEDIUM", "Expected results need execution or external evidence",
        f"{len(not_message)} of {len(scenarios)} expected results are device, host, issuer or settlement outcomes "
        "that message-level validation cannot confirm. The workbook has no test data, actual-result or evidence columns.",
        [s["id"] for s in not_message], ARTIFACT_OWNER, "Add test data references, actual results and evidence per scenario."))
    counts = Counter(item["check"] for item in parsed["structure"])
    out.append(finding(
        "MCV-A06", "AI_ARTIFACT_DEFECT", "LOW", "Structure: no line IDs, stray whitespace, several checks in one cell",
        "; ".join(f"{k}: {v}" for k, v in sorted(counts.items())) + ".",
        list(OrderedDict.fromkeys(item["line"] for item in parsed["structure"] if item["check"] != "NO_VALIDATION_LINE_IDS")),
        ARTIFACT_OWNER, "Give each validation an ID, trim whitespace and keep one check per cell."))
    a07 = curated_findings["MCV-A07"]
    out.append(finding("MCV-A07", a07["class"], a07["severity"], a07["title"], a07["detail"], a07["affected"],
                       a07["owner"], a07["disposition"]))
    out.append(finding(
        "MCV-T01", "TEST_SOLUTION_GAP", "MEDIUM", "Source-supported checks without a trained Test Solution chain",
        "; ".join(f"{line['id']} ({line['coverage']}"
                  + (": " + ", ".join(d["id"] for d in line["draftBrs"]) if line["draftBrs"] else "") + ")"
                  for line in gaps) + ".",
        [line["id"] for line in gaps], TEST_SOLUTION_OWNER,
        "Derive canonical BR -> TS -> TC -> TD chains (promote the Section 10.1 drafts where they exist)."))
    out.append(finding(
        "MCV-T02", "CHAIN_GAP", "MEDIUM", "Cited Test Solution BRs that are still in the Mastercard review queue",
        "; ".join(f"{b}: {review_queue[b]['reasons'][0]}" for b in queued) + ".",
        [line["id"] for line in chain_lines], TEST_SOLUTION_OWNER,
        "Close the review-queue items (anchor or split the shared scenarios) and retrain the context."))
    out.append(finding(
        "MCV-T03", "TEST_SOLUTION_GAP", "MEDIUM" if multi_segment else "LOW",
        "Multi-segment anchors are excluded from the Mastercard context scope",
        f"The context scope compares the whole segment string, so anchors such as '100,111' never match (handbook L22). "
        f"{len(multi_segment)} BRs would enter scope if segments were split on ','"
        + (": " + ", ".join(item["id"] for item in multi_segment) if multi_segment else "") + ".",
        [item["id"] for item in multi_segment], TEST_SOLUTION_OWNER,
        "Split segment strings on ',' in in_mc_scope, then retrain and re-run this validation.",
        "OPEN" if multi_segment else "RESOLVED"))
    t04 = curated_findings["MCV-T04"]
    open_t04 = "BR-APPG-CODE-3" in mc.TEMPLATE_MAP["TS-020"][0]
    out.append(finding("MCV-T04", t04["class"], t04["severity"], t04["title"], t04["detail"], t04["affected"],
                       t04["owner"], t04["disposition"], "OPEN" if open_t04 else "RESOLVED"))
    reference = ba["reference"]
    out.append(finding(
        "MCV-T05", "PROVENANCE_OR_VERSION_GAP", "LOW", "The BA reference review is not bound to this workbook",
        f"The BA review was transcribed from a screenshot (originalFileInRepository={reference['originalFileInRepository']}, "
        f"originalFileHash={reference['originalFileHash']}) although the workbook is in the repository. "
        f"{len(transcription)} transcription differences were verified ("
        + ", ".join(item["ts"] for item in transcription) + "); none changes a disposition. "
        f"Automated name/FR/element fidelity mismatches: {len(fidelity)}.",
        [item["ts"] for item in transcription], TEST_SOLUTION_OWNER,
        "Bind the BA review to the workbook hash and correct the four transcribed claims."))
    return out


# ---------------------------------------------------------------- rendering

FILLS = {"CONTRADICTED": "FFC7CE", "SUPPORTED_WITH_CORRECTION": "FFEB9C", "EXTERNAL_POLICY": "DDEBF7",
         "NOT_TESTABLE_AS_WRITTEN": "E7E6E6", "SUPPORTED": "C6EFCE", "REJECTED_NOT_INTAKE_READY": "FFC7CE",
         "REVIEW_REQUIRED": "DDEBF7", "ACCEPTED_WITH_GAPS": "FFEB9C", "ACCEPTED": "C6EFCE", "MISMATCH": "FFC7CE",
         "MATCH": "C6EFCE", "HIGH": "FFC7CE", "MEDIUM": "FFEB9C", "LOW": "F2F2F2"}


def csv_text(rows):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    for row in rows:
        writer.writerow({key: ("'" + str(value) if str(value).startswith(("=", "+", "-", "@")) else str(value))
                         for key, value in row.items()})
    return stream.getvalue()


def evidence_text(line):
    parts = []
    for item in line["evidence"]:
        if "absence" in item:
            absence = item["absence"]
            parts.append(f"{absence['section']} (lines {absence['lines']}): no '{absence['term']}'")
        else:
            parts.append(f"{item['section']} ({Path(item['path']).name}:{item['line']})")
    return "; ".join(parts)


def line_rows(report):
    scenarios = {s["id"]: s for s in report["scenarios"]}
    rows = []
    for line in report["lines"]:
        scenario = scenarios[line["ts"]]
        rows.append(OrderedDict([
            ("Line ID", line["id"]), ("FR ID", scenario["fr"]), ("TS ID", scenario["id"]),
            ("Test Scenario", scenario["name"]), ("Validation (as written)", line["text"]), ("Verdict", line["verdict"]),
            ("Severity", line["severity"]), ("Finding Classes", ";".join(line["classes"])),
            ("Assessment", line["assessment"]), ("Corrected Requirement", line["correction"]),
            ("ATL105 Evidence", evidence_text(line)),
            ("Test Solution BRs", "; ".join(f"{b['id']} [{b['state']}]" for b in line["testSolutionBrs"])),
            ("Section 10.1 Draft BRs", "; ".join(d["id"] for d in line["draftBrs"])),
            ("Related BRs (context only)", "; ".join(f"{b['id']} [{b['state']}]" for b in line["relatedBrs"])),
            ("Test Solution Coverage", line["coverage"]), ("BA Lessons", ";".join(line["lessons"])),
            ("Open Questions", ";".join(line["questions"]))]))
    return rows


def scenario_rows(report):
    return [OrderedDict([
        ("TS ID", s["id"]), ("FR ID", s["fr"]), ("Test Scenario", s["name"]), ("Acceptance Criteria", s["acceptanceCriteria"]),
        ("Outcome", s["outcome"]), ("Worst Line Verdict", s["worstLineVerdict"]), ("Lines", ", ".join(s["lineIds"])),
        ("Expected Result", s["expectedResult"]), ("Expected Result Verdict", s["expectedVerdict"]),
        ("Expected Result Note", s["expectedNote"]), ("BA Review Disposition", s["baDisposition"] or ""),
        ("BA Corrected Requirement", s["baCorrectedRequirement"] or ""),
        ("Template Coverage (Test Solution)", s["templateCoverage"] or ""), ("Training State", s["trainingState"] or ""),
        ("Mapped Test Solution BRs", "; ".join(s["mappedTestSolutionBrIds"])), ("Open Questions", ";".join(s["questions"]))])
        for s in report["scenarios"]]


def finding_rows(report):
    return [OrderedDict([("ID", f["id"]), ("Class", f["class"]), ("Severity", f["severity"]), ("Status", f["status"]),
                         ("Title", f["title"]), ("Detail", f["detail"]), ("Affected", "; ".join(f["affected"])),
                         ("Owner", f["owner"]), ("Disposition", f["disposition"])]) for f in report["findings"]]


def element_rows(report):
    return [OrderedDict([("Line ID", e["line"]), ("Validation (as written)", e["text"]), ("Intended Subject", e["subject"]),
                         ("Cited Element", e["citedElement"]), ("Cited Element Name (Ch. 13)", e["citedName"]),
                         ("Spec Line", e["citedSpecLine"]), ("Correct Element", e["correctElement"]),
                         ("Correct Element Name (Ch. 13)", e["correctName"]), ("Correct Spec Line", e["correctSpecLine"]),
                         ("Result", e["result"])]) for e in report["elementCheck"]]


def transcription_rows(report):
    return [OrderedDict([("TS ID", t["ts"]), ("Workbook Text", t["workbookText"]),
                         ("BA Review Transcription", t["baClaimExcerpt"] or "(not transcribed)"),
                         ("Difference", t["difference"]), ("Impact", t["impact"]),
                         ("BA Reference Claim", t["baReferenceClaim"])]) for t in report["transcriptionDifferences"]]


def md_table(header, rows):
    return ("| " + " | ".join(header) + " |\n|" + "|".join("---" for _ in header) + "|\n"
            + "".join("| " + " | ".join(str(cell).replace("|", "\\|").replace("\n", " ") for cell in row) + " |\n"
                      for row in rows))


def readme(report):
    summary = report["summary"]
    lines = report["lines"]
    verdicts = Counter(line["verdict"] for line in lines)
    contradicted = [line for line in lines if line["verdict"] == "CONTRADICTED"]
    trained = sum(line["coverage"] == "TRAINED_CHAIN" for line in lines)
    message_level = sum(s["expectedVerdict"] == "MESSAGE_LEVEL" for s in report["scenarios"])
    fr_rows = []
    for fr in report["functionalRequirements"]:
        fr_lines = [line for line in lines if line["ts"] in fr["scenarios"]]
        counts = Counter(line["verdict"] for line in fr_lines)
        fr_rows.append([fr["id"], fr["name"], len(fr["scenarios"]), len(fr_lines)] + [counts.get(v, 0) for v in VERDICTS])
    md = [
        "# Mastercard Credit Purchase FR Matrix: Workbook Validation\n",
        f"This review validates `{report['workbook']['path']}` (sheet `{report['workbook']['sheet']}`, SHA-256 "
        f"`{report['workbook']['sha256']}`) line by line against the current Test Solution for ATL105 "
        f"{report['specificationVersion']}: the ATL105 source text, the canonical Test Solution BRs and their Mastercard "
        "training status, the Section 10.1 credit-card draft BRs, and the Mastercard BA reference review. Status is "
        f"{report['status']}; no SME approval or execution certification is implied.\n",
        f"## Outcome: {report['outcome']}\n",
        f"As written, the workbook cannot be used as a test oracle: {verdicts['CONTRADICTED']} validation lines point at the "
        f"wrong ATL105 field and {verdicts['NOT_TESTABLE_AS_WRITTEN']} name no testable field or rule, so tests derived "
        "from them would check the wrong thing. The BR-01 / FR / scenario structure is sound and is already the business "
        "layer of the Test Solution FR matrix. Every line's corrected requirement is in `line-validation.csv` and "
        "`validation.xlsx`.\n",
        "## Summary\n",
        md_table(["Measure", "Result"], [
            ["Business requirement / FRs / scenarios / validation lines",
             f"1 / {summary['functionalRequirements']} / {summary['scenarios']} / {summary['validationLines']}"],
            ["Lines supported as written", verdicts["SUPPORTED"]],
            ["Lines supported once corrected", verdicts["SUPPORTED_WITH_CORRECTION"]],
            ["Lines contradicted by ATL105", verdicts["CONTRADICTED"]],
            ["Lines that depend on contracts outside ATL105", verdicts["EXTERNAL_POLICY"]],
            ["Lines not testable as written", verdicts["NOT_TESTABLE_AS_WRITTEN"]],
            ["Lines backed by a trained Test Solution chain", f"{trained} of {len(lines)}"],
            ["Element citations naming the wrong element", f"{summary['elementMismatches']} of {summary['elementCitations']}"],
            ["Expected results confirmable from the message alone", f"{message_level} of {summary['expectedResults']}"],
        ]),
        "\n## Scenario Outcomes\n",
        md_table(["Outcome", "Scenarios"], [[o, ", ".join(s["id"] for s in report["scenarios"] if s["outcome"] == o) or "-"]
                                            for o in OUTCOMES]),
        "\n## Wrong Field Mappings\n",
        md_table(["Line", "As written", "Correct requirement", "ATL105 evidence"],
                 [[line["id"], line["text"].strip(), line["correction"], evidence_text(line)] for line in contradicted]),
        "\n## Element Citation Check (Chapter 13)\n",
        md_table(["Line", "Cited", "Cited element is", "Intended", "Correct element", "Result"],
                 [[e["line"], f"Element {e['citedElement']}", e["citedName"], e["subject"],
                   f"Element {e['correctElement']} {e['correctName']}", e["result"]] for e in report["elementCheck"]]),
        "\n## Line Verdicts by FR\n",
        md_table(["FR", "Name", "Scenarios", "Lines", "Supported", "With correction", "Contradicted", "External", "Not testable"],
                 fr_rows),
        "\n## Findings\n",
        md_table(["ID", "Class", "Severity", "Status", "Title", "Owner"],
                 [[f["id"], f["class"], f["severity"], f["status"], f["title"], f["owner"]] for f in report["findings"]]),
        "\nFull detail, affected IDs and dispositions are in `findings.csv`.\n",
        "\n## BA Reference Review Transcription Differences\n",
        "The existing [BA reference review](../mastercard-ba-reference-review-v1/README.md) was transcribed from a "
        "screenshot. Checked against the actual workbook:\n\n",
        md_table(["TS", "Workbook", "BA review transcription", "Impact"],
                 [[t["ts"], t["workbookText"], t["baClaimExcerpt"] or "(not transcribed)", t["impact"]]
                  for t in report["transcriptionDifferences"]]),
        "\n## Deliverables\n",
        "- `validation.xlsx`: line validation, scenario outcomes, findings, element check and transcription sheets.\n"
        "- `line-validation.csv`: one row per validation line with verdict, correction, evidence and Test Solution status.\n"
        "- `findings.csv`: findings with class, severity, owner and disposition.\n"
        "- `validation.json`: the complete machine-readable review, including input hashes.\n",
        "\n## Re-run\n",
        "From the validator module:\n\n```powershell\n"
        "python scripts/validate-mastercard-fr-matrix.py           # regenerate this folder\n"
        "python scripts/validate-mastercard-fr-matrix.py --check   # fail if these outputs are stale\n"
        "python -m pytest scripts/test_validate_mastercard_fr_matrix.py\n```\n\n"
        "The script refuses to run if the workbook's SHA-256, any validation line or expected result, an element "
        "definition, a source anchor or a cited BR no longer matches, so a changed workbook must be re-reviewed.\n",
        "\n## Limitations\n",
        "- Line assessments are curated Test Solution judgements; they remain DRAFT_REVIEW_REQUIRED until SME/TBA review.\n"
        "- Section 10.1 draft BRs are cited as source-derived evidence; they are not trained chains.\n"
        "- External contracts (BIN/brand, issuer status, merchant configuration, routing, clearing, settlement, audit) "
        "are not available, so those checks stay REVIEW_REQUIRED (questions MC-Q01 to MC-Q09, register SEG100-SME-003).\n"
        "- Nothing was executed; expected results that need host, issuer or device evidence are not confirmed.\n",
    ]
    return "\n".join(md)


def write_sheet(ws, rows, widths, color_columns=()):
    header = list(rows[0])
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row[name] for name in header])
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(header) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        for name in color_columns:
            cell = ws.cell(r, header.index(name) + 1)
            if cell.value in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def workbook_bytes(report):
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    rows = [["Mastercard Credit Purchase FR Matrix: Workbook Validation", ""],
            ["Workbook", report["workbook"]["path"]], ["SHA-256", report["workbook"]["sha256"]],
            ["Specification", f"{report['specification']} {report['specificationVersion']}"],
            ["Status", report["status"]], ["Outcome", report["outcome"]], ["Reviewed on", report["reviewedOn"]], ["", ""]]
    for key, value in report["summary"].items():
        rows.append([key, json.dumps(value) if isinstance(value, dict) else value])
    for row in rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True, size=14)
    summary["B6"].fill = PatternFill("solid", fgColor=FILLS.get(report["outcome"], "FFFFFF"))
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 110
    write_sheet(wb.create_sheet("Line Validation"), line_rows(report),
                [11, 8, 8, 26, 34, 24, 9, 30, 60, 50, 42, 52, 30, 30, 20, 14, 14], ("Verdict", "Severity"))
    write_sheet(wb.create_sheet("Scenarios"), scenario_rows(report),
                [8, 8, 28, 34, 24, 24, 22, 28, 28, 40, 24, 50, 20, 22, 44, 16], ("Outcome", "Worst Line Verdict"))
    write_sheet(wb.create_sheet("Findings"), finding_rows(report), [9, 24, 9, 10, 40, 90, 50, 26, 50], ("Severity",))
    write_sheet(wb.create_sheet("Element Check"), element_rows(report), [11, 34, 30, 10, 26, 10, 10, 26, 12, 11], ("Result",))
    write_sheet(wb.create_sheet("BA Transcription"), transcription_rows(report), [8, 40, 40, 50, 40, 60])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def build(workbook_path=None):
    report = build_report(workbook_path)
    outputs = OrderedDict([("README.md", readme(report)),
                           ("validation.json", json.dumps(report, indent=2, ensure_ascii=True) + "\n"),
                           ("line-validation.csv", csv_text(line_rows(report))),
                           ("findings.csv", csv_text(finding_rows(report)))])
    return report, outputs, workbook_bytes(report)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Validate the Mastercard Credit Purchase FR matrix workbook.")
    parser.add_argument("--check", action="store_true", help="fail if the written outputs are stale; write nothing")
    args = parser.parse_args(argv)
    report, outputs, xlsx = build()
    if args.check:
        stale = [name for name, text in outputs.items()
                 if not os.path.exists(long_path(OUT_DIR / name)) or read_bytes(OUT_DIR / name).decode("utf-8") != text]
        print("Stale validation outputs: " + ", ".join(stale) if stale else "Validation outputs are current.")
        return 1 if stale else 0
    os.makedirs(long_path(OUT_DIR), exist_ok=True)
    for name, text in outputs.items():
        with open(long_path(OUT_DIR / name), "w", encoding="utf-8", newline="") as handle:
            handle.write(text)
    with open(long_path(OUT_DIR / "validation.xlsx"), "wb") as handle:
        handle.write(xlsx)
    summary = report["summary"]
    print(f"Outcome {report['outcome']} | scenarios {summary['scenarios']} | lines {summary['validationLines']} "
          f"{summary['lineVerdicts']} | scenario outcomes {summary['scenarioOutcomes']} | element mismatches "
          f"{summary['elementMismatches']}/{summary['elementCitations']} | findings {summary['findings']}")
    print(f"Wrote {rel(OUT_DIR)}/ ({', '.join(list(outputs) + ['validation.xlsx'])})")
    return 0


if __name__ == "__main__":
    sys.exit(main())

import argparse
import hashlib
import json
import os
import re
from html import escape
from pathlib import Path


MODULE = Path(__file__).resolve().parents[1]
ARTIFACTS = {"BR": "businessRequirements", "TS": "testScenarios", "TC": "testCases", "TD": "testData"}


def read(file):
    return json.loads(file.read_text(encoding="utf-8-sig"))


def digest(file):
    with file.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def counts(package):
    result = {}
    for label, key in ARTIFACTS.items():
        rows = package.get(key)
        if not isinstance(rows, list):
            raise ValueError("Missing canonical artifact array: " + key)
        ids = [row.get("id") for row in rows]
        if len(set(ids)) != len(ids) or any(not value for value in ids):
            raise ValueError("Missing/duplicate artifact IDs: " + key)
        result[label] = len(rows)
    return result


def native(file):
    normalized = Path(os.path.normpath(str(file)))
    return Path("\\\\?\\" + str(normalized.resolve())) if os.name == "nt" and not str(normalized).startswith("\\\\?\\") else normalized


def gather(review):
    independent = MODULE / "specifications/ATL105/test-output/test-solution-independent-review"
    run_file = review / "complete-handoff-analysis.json"
    coverage_file = review / "coverage-view-assessment.json"
    training_file = independent / "all-test-solution-br-ts-tc-td-training-package.json"
    complete_file = independent / "complete-all-test-solution-br-ts-tc-td-package.json"
    approved_file = independent / "approved-br-ts-tc-td-chain-package.json"
    crosswalk_file = independent / "atomic-br-ai-crosswalk.json"
    matrix_file = review / "late-traceability-matrix-assessment.json"
    run = read(run_file)
    coverage = read(coverage_file)
    training = read(training_file)
    complete = read(complete_file)
    approved = read(approved_file)
    historical = read(crosswalk_file)
    matrix = read(matrix_file)
    result = {"asOf": "2026-10-07", "delivery": "ATL105 2026-3 | October 5 Run1",
              "ai": {"BR": run["businessRequirements"], "TS": run["scenarios"], "TC": run["caseReconciliation"]["records"],
                     "logicalCasesWithTD": run["producerPayloadTraceability"]["totalTestCases"],
                     "physicalTDRequestFiles": run["producerPayloadTraceability"]["totalDataFiles"],
                     "physicalRequestMetadataPairs": run["producerPayloadTraceability"]["totalDataFiles"]},
              "testSolutionPopulatedAggregate": counts(training), "testSolutionStructuralPackage": counts(complete),
              "testSolutionApprovedChainPackage": counts(approved), "independentCatalogRuleInventory": coverage["catalogRuleInventory"],
              "independentCatalogSegments": coverage["catalogSegments"],
              "structuralLinkage": {"brWithTc": run["caseReconciliation"]["linkedDistinctRequirements"],
                  "scenarioWithTc": run["caseReconciliation"]["linkedDistinctScenarios"], "scenarioWithoutTc": run["caseReconciliation"]["scenariosWithoutCase"]},
              "perfectFullChainMatchCount": "NOT_ESTABLISHED_FOR_OCTOBER_DELIVERY", "executionCertified": False,
              "historicalCrosswalkExcluded": {"source": historical["aiSource"], "counts": historical["counts"], "reason": "Different delivery; candidate mappings are not confirmed full-chain matches"},
              "matrix": {"aiDeclaredRows": matrix["reconstruction"]["producerFlatRowsDeclared"],
                         "independentReconstructedRows": matrix["reconstruction"]["leafRows"],
                         "displayedFlatRows": matrix["independent"]["leafRows"], "displayedBrDetails": matrix["independent"]["detailRequirementCount"],
                         "detailedStatusesAgree": matrix["reconstruction"]["detailStatusAgreementCount"],
                         "leafStatusDifferences": len(matrix["reconstruction"]["flatExcerptStatusDisagreements"])},
              "inputSha256": {str(file.relative_to(MODULE)).replace("\\", "/"): digest(file) for file in [run_file, coverage_file, training_file, complete_file, approved_file, crosswalk_file, matrix_file]}}
    result["structuralLinkage"]["brPercent"] = round(100 * result["structuralLinkage"]["brWithTc"] / result["ai"]["BR"], 2)
    if any(result["testSolutionApprovedChainPackage"].values()):
        raise ValueError("Approved-package evidence changed; reassess presentation match claims")
    return result


def examples(review):
    import ijson
    training_file = MODULE / "specifications/ATL105/test-output/test-solution-independent-review/all-test-solution-br-ts-tc-td-training-package.json"
    training = read(training_file)
    run = read(review / "complete-handoff-analysis.json")
    archive = Path(run["archive"])
    catalogs = archive / "scenerio_req_5_oct/scenerio_req_5_oct"
    br_file = catalogs / "step5_requirements/approved/requirement_catalog.json"
    scenario_file = catalogs / "scenarios/approved/approved_scenarios.json"
    case_file = archive / "pipeline_run_artifacts/test_case_candidates.json"
    brs = {row["id"]: row for row in read(br_file)["requirements"]}
    scenarios = {row["id"]: row for row in read(scenario_file)["scenarios"]}
    selected = None
    with case_file.open("rb") as stream:
        for case in ijson.items(stream, "test_cases.item", use_float=True):
            if case["id"] == "TC-001436":
                selected = case
                break
    if selected is None:
        raise ValueError("Pinned AI example is not in the current delivery")
    payload_file = archive / "pipeline_run_artifacts/qe_shaped_test_data/TC-001436.json"
    payload = read(payload_file)
    segment = payload["Financial Transaction Request"]["Standard Segment"]
    requirement_ids = set(selected.get("requirement_ids") or []) | {selected.get("requirement_id")}
    evidence = []
    choices = [
        ("Sequence format", "REQ-SRC-ATL105-PDF-001:097", "BR-SEG100-E86-001", "SCN-SEG100-E86-001", "TC-SEG100-E86-001", "TD-SEG100-E86-001", "SequenceNumber", "sequenceNumber", "FIELD_PREDICATE_ALIGNMENT_ONLY"),
        ("Segment identity", "REQ-SRC-ATL105-PDF-001:095", "BR-SEG100-ID-001", "SCN-SEG100-ID-001", "TC-SEG100-ID-001", "TD-SEG100-ID-001", "SegmentType", "segmentType", "FIELD_VALUE_ALIGNMENT_ONLY"),
        ("Partial approval context", "REQ-SRC-ATL105-PDF-001:142", "BR-SEG100-E121-002", "SCN-SEG100-E121-002", "TC-SEG100-E121-003", "TD-SEG100-E121-003", "PartialApprovalIndicator", "partialApprovalIndicator", "NOT_EQUIVALENT_CONTEXT_OBLIGATION_NOT_PROVEN"),
    ]
    by_kind = {key: {row["id"]: row for row in training[key]} for key in ARTIFACTS.values()}
    for title, ai_br, test_br, test_ts, test_tc, test_td, ai_field, test_field, disposition in choices:
        if ai_br not in requirement_ids or ai_br not in scenarios[selected["scenario_id"]].get("requirement_ids", [scenarios[selected["scenario_id"]].get("requirement_id")]):
            raise ValueError("AI example lacks expected BR/TS/TC linkage: " + ai_br)
        requirement = by_kind["businessRequirements"][test_br]
        scenario = by_kind["testScenarios"][test_ts]
        case = by_kind["testCases"][test_tc]
        data = by_kind["testData"][test_td]
        if test_br not in scenario["requirementIds"] or test_ts not in case["scenarioIds"] or test_tc not in data["testCaseIds"]:
            raise ValueError("Independent example chain is broken")
        actual = segment[ai_field]
        independent = data["payload"]["request"]["dataSection2"]["standardSegment"][test_field]
        if ai_field == "SequenceNumber" and not all(re.fullmatch("[0-9]{6}", value) for value in [actual, independent]):
            raise ValueError("Sequence-format example no longer proves the stated predicate")
        if ai_field == "SegmentType" and (actual != "100" or independent != "100"):
            raise ValueError("Identity example no longer proves the stated field equality")
        evidence.append({"title": title, "disposition": disposition,
            "ai": {"BR": ai_br, "brStatement": brs[ai_br].get("statement"), "TS": selected["scenario_id"],
                   "scenarioClaim": scenarios[selected["scenario_id"]].get("name"), "TC": selected["id"],
                   "TD": payload_file.name, "field": ai_field, "value": actual, "expectedResponse": selected.get("expected_response")},
            "test": {"BR": test_br, "brStatement": requirement.get("title"), "TS": test_ts, "TC": test_tc, "TD": test_td,
                     "sourceAnchors": requirement["sourceAnchors"], "field": test_field, "value": independent,
                     "expectedValidation": data.get("expectedValidation"), "readiness": data.get("readiness")},
            "wholeChainConfirmed": False})
    return {"examples": evidence, "inputSha256": {str(file): digest(file) for file in [training_file, br_file, scenario_file, case_file, payload_file]},
            "scope": "Actual producer/test artifact comparisons; values and links verified, full behavioral equivalence/approval not granted"}


def slides(facts, proof):
    ai = facts["ai"]
    populated = facts["testSolutionPopulatedAggregate"]
    structural = facts["testSolutionStructuralPackage"]
    sequence, identity, partial = proof["examples"]
    return [
        {"title": "ATL105 | AI validation business briefing", "subtitle": "Fiserv | October 7, 2026 | October 5 Run1 delivery",
         "lines": ["AI produced substantial BR, scenario, case and payload inventories.", "Independent reconstruction makes the full structural graph auditable.", "Structural links and matching field values do not establish full business equivalence or host acceptance.", "Decision today: review the discrepancies and authorize a defined evidence-qualified scope."],
         "notes": "Lead with progress, then explain that inventory, semantic matching and execution certification are different measures. Do not describe 53.13% as validated semantic coverage."},
        {"title": "1 | What did the AI Solution create?", "rows": [["Artifact", "Count"], ["BR", f"{ai['BR']:,}"], ["TS", f"{ai['TS']:,}"], ["TC", f"{ai['TC']:,}"], ["TD request files", f"{ai['physicalTDRequestFiles']:,}"], ["Logical TCs with physical TD", f"{ai['logicalCasesWithTD']:,} / {ai['TC']:,}"]],
         "lines": ["Each of the 21,210 request files has a metadata companion; metadata is not another TD.", "Flows have multiple payload legs. Therefore physical TD files and logical TC counts differ.", "27 TCs have no physical TD; 3,732 scenarios have no TC."], "notes": "Say 21,210 physical request-data files covering 21,096 logical cases. Do not say 42,420 test data artifacts by counting metadata twice."},
        {"title": "2 | Is there a complete AI traceability matrix?", "lines": ["Yes, a matrix document was delivered; its displayed content is truncated.", "AI declares 61,107 rows, but Markdown shows 2,000 rows and 200 BR details.", "Our independent reconstruction has 61,044 leaves and represents all 6,887 BRs.", "3,659 BRs have TC links (53.13% internal linkage); 8,947 scenarios have TCs.", "Producer's 3,661 linked BRs and 92 orphan scenarios differ from frozen 3,659 and 155.", "All 200 displayed BR-level statuses agree; 75 leaf labels inherit a BR's FULLY_TRACED status."],
         "notes": "A matrix exists, but not all BRs have complete generated chains. The 63 additional producer rows are unconfirmed without the omitted producer edges. Reconstruction is independent evidence, not the missing producer JSON."},
        {"title": "3 | What has the Test Solution identified?", "rows": [["Inventory", "BR", "TS", "TC", "TD"], ["Populated aggregate", *[f"{populated[key]:,}" for key in ARTIFACTS]], ["Structural package", *[f"{structural[key]:,}" for key in ARTIFACTS]], ["Approved-chain package", "0", "0", "0", "0"]],
         "lines": ["Independent segment rule-catalog inventory: 601 rules across 49 segments.", "Structural package includes placeholders: 618 TS, 619 TC and 641 TD.", "These are inventory counts, not 1:1 AI matches or execution-ready chains."],
         "notes": "Use populated aggregate as the primary answer and explain structural inventory separately. Neither is an approved semantic denominator for the October AI delivery. The approved-chain package is empty, not proof that no source rules exist."},
        {"title": "4 | How are AI and Test artifacts matched?", "lines": ["Freeze delivery/version and source hashes; preserve original producer artifacts.", "Normalize formats through adapters; preserve each producer's local IDs.", "Match source anchors plus transaction/card/network/lifecycle context.", "Compare BR conditions, obligations, exceptions and expected behavior.", "Trace TS objective > TC assertions > TD predicates and dependencies.", "Assign CONFIRMED, REVIEW_REQUIRED or MISSING with evidence and owners."],
         "notes": "Different IDs and legal data values can match. Literal constants must match; variable data may differ if both satisfy the same rule and preserve dependencies. Similar titles, shared IDs and AI confidence are only candidate evidence."},
        {"title": "5 | How many perfect matches can we claim?", "lines": ["No evidence-qualified full-chain perfect-match count is established for October Run1.", "Zero confirmed full-chain matches can currently be claimed; this is not proof of zero real equivalence.", "200/200 BR status agreement is status reconciliation, not business-rule equivalence.", "40/40 synthetic probes passed two scoped predicates; they are not forty approved AI tests.", "The 522 atomic candidates belong to September Run2 and are not October confirmed matches."],
         "notes": "Answer honestly: current confirmed count is none established. We have concrete field-level alignment and partial candidates, not an approved complete cross-producer matrix."},
        {"title": "6A | Real example: different data, same predicate", "rows": [["Layer", "AI Solution", "Independent Test Solution"], *[[key, sequence['ai'][key], sequence['test'][key]] for key in ['BR', 'TS', 'TC', 'TD']], ["Sequence value", sequence['ai']['value'], sequence['test']['value']]],
         "lines": ["Both actual values satisfy exactly six ASCII digits.", "Disposition: verified field-predicate alignment; partial full-chain comparison candidate.", "AI expected response is empty; independent fixture readiness is not declared. Range/allocation/lifecycle and complete outcomes remain unproven."],
         "notes": "Do not call this a perfect BR match. AI BR also carries range/application language, while the Test BR shown is a six-digit representation obligation. Show the full statements in the companion brief."},
        {"title": "6B | Real example: identity aligns, context differs", "lines": [f"Identity: AI {identity['ai']['TC']} and Test {identity['test']['TC']} both carry Segment Type 100.", f"AI identity BR {identity['ai']['BR']} is generic; Test {identity['test']['BR']} explicitly fixes Standard Segment type to 100.", f"Partial approval: AI value {partial['ai']['value']} vs Test value {partial['test']['value']}.", "The Test BR requires value 5 in Amex prepaid balance-receipt context.", "AI value 0 does not prove that value-5 context. Different values are not automatically interchangeable.", "These examples prove scoped alignment or a context gap, not approved full-chain equivalence."],
         "notes": "The partial-approval example is a not-proven context comparison, not an automatic AI defect. A positive value-0 case may legitimately test a different branch."},
        {"title": "7 | How can we filter artifacts for analysis?", "lines": ["Run report: no-TC search, generation cause, producer verdict and evidence conditions; selected CSV export.", "Coverage view: overall, segment, transaction target, response-code declaration and scenario-type summaries.", "Semantic report: 100-case intent/condition filters, artifact/BR search, detail and CSV export.", "Code-1 queue: producer transaction label, source/context filters and paginated export.", "Full reconstruction CSV: filter BR, scenario, TC, leaf status and TD paths in Excel.", "Scope warning: no-TC view covers 3,732 scenarios; semantic case view covers 100, not the entire delivery."],
         "notes": "Demo one scoped filter at a time and state its denominator. Coverage selectors are summary dimensions, not a full 21,123-case explorer. Unattributed cases remain separately in the full JSON."},
        {"title": "Business decision and next evidence", "lines": ["Request the complete producer matrix and source catalog hashes/changed artifacts.", "Resolve the 63 additional rows and BR :521 / :3226 mappings without inventing edges.", "Adjudicate a narrow business scope against independent source-derived requirements.", "Supply approved control fixtures, actual wire, authoritative host expectations and a permitted test environment.", "Track structural, semantic and executed coverage separately; retain every assessment history point."],
         "notes": "Close with the specific evidence and authorization needed. Source/SME review and processor authorization are business controls, not issues to bypass for a green report."},
    ]


def generate_outputs(output, facts, proof, review, html_only=False):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from PIL import ImageFont
    content = slides(facts, proof)
    markdown = ["# Fiserv Business Brief | ATL105 AI Validation", "", "As of October 7, 2026. Scope: October 5 Run1; figures are inventories unless explicitly stated otherwise.", ""]
    html = ["<!doctype html><html lang=\"en\"><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width,initial-scale=1\"><title>Fiserv | ATL105 Business Brief</title><style>body{margin:0;color:#202526;background:#f4f6f5;font:16px 'Segoe UI',sans-serif;letter-spacing:0}header{background:white;border-top:5px solid #07665e;padding:24px}main{max-width:1100px;margin:auto;padding:24px}h1{font:700 28px Georgia,serif}h2{font-size:21px;color:#07665e}section{padding:24px 0;border-bottom:1px solid #cdd7d5}table{border-collapse:collapse;width:100%;background:white}th,td{padding:10px;border:1px solid #cdd7d5;text-align:left;overflow-wrap:anywhere}th{background:#e5efed}li,p{line-height:1.55}a{color:#07665e}.table-wrap{overflow-x:auto}.scope{color:#626b6c}code{overflow-wrap:anywhere}@media(max-width:600px){main,header{padding:16px}table{min-width:650px}}@media print{body{background:white}section{break-before:page}header{break-after:page}main{padding:0}}</style><header><h1>ATL105 AI Validation | Business Brief</h1><p>Fiserv | October 7, 2026 | October 5 Run1</p><p><a href=\"FISERV-BUSINESS-BRIEF.pptx\">Presentation with speaker notes</a> | <a href=\"FISERV-BUSINESS-BRIEF.pdf\">Brief PDF</a> | <a href=\"../AI-ARTIFACT-FILTER-VIEW.html\">Run report / filters</a> | <a href=\"../SEMANTIC-FIRST-BATCH-REPORT.html\">Semantic report / examples</a></p></header><main>"]
    pdf_story = []
    heading = ParagraphStyle("Heading", fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=colors.HexColor("#07665e"), spaceAfter=16)
    body = ParagraphStyle("Body", fontName="Helvetica", fontSize=11, leading=16, spaceAfter=9)
    small = ParagraphStyle("Cell", parent=body, fontSize=9, leading=12)
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    def box(slide, left, top, width, height, text, size=20, bold=False, color="202526"):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        usable_width = width * 72 - 15
        usable_height = height * 72 - 8
        def estimated_height(font_size):
            font = ImageFont.truetype("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf", font_size)
            total = 0
            for paragraph in text.split("\n"):
                line = ""
                lines = 1
                for word in paragraph.split():
                    candidate = (line + " " + word).strip()
                    if line and font.getlength(candidate) > usable_width:
                        lines += 1
                        line = word
                    else:
                        line = candidate
                total += lines * font_size * 1.25 + 8
            return total
        while size > 14 and estimated_height(size) > usable_height:
            size -= 1
        if estimated_height(size) > usable_height:
            raise ValueError("Slide textbox exceeds available height: " + text[:60])
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(8)
        return shape
    for index, item in enumerate(content):
        title = item["title"]
        markdown.extend(["## " + title, ""])
        html.append("<section><h2>" + escape(title) + "</h2>")
        if index:
            pdf_story.append(PageBreak())
        pdf_story.append(Paragraph(escape(title), heading))
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box(slide, 0.6, 0.35, 12.1, 0.85, title, 28, True, "07665E")
        top = 1.45
        if item.get("rows"):
            rows = item["rows"]
            markdown.extend(["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"])
            markdown.extend("| " + " | ".join(map(str, row)) + " |" for row in rows[1:])
            markdown.append("")
            html.append('<div class="table-wrap"><table>' + "".join("<tr>" + "".join("<td>" + escape(str(cell)) + "</td>" for cell in row) + "</tr>" for row in rows) + "</table></div>")
            table = Table([[Paragraph(escape(str(cell)), small) for cell in row] for row in rows], colWidths=[730 / len(rows[0])] * len(rows[0]))
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5efed")), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cdd7d5")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            pdf_story.extend([table, Spacer(1, 16)])
            height = min(3.35, .47 * len(rows))
            table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.6), Inches(top), Inches(12.1), Inches(height))
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table_shape.table.cell(row_index, column_index)
                    cell.text = str(value)
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Aptos"
                        paragraph.font.size = Pt(15 if len(rows[0]) <= 3 else 17)
                        paragraph.font.color.rgb = RGBColor.from_string("202526")
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("E5EFED" if row_index == 0 else "FFFFFF")
            top += height + .25
        lines = item["lines"]
        markdown.extend("- " + line for line in lines)
        markdown.extend(["", "**Speaker note:** " + item["notes"], ""])
        html.append("<ul>" + "".join("<li>" + escape(line) + "</li>" for line in lines) + "</ul><p class=\"scope\"><b>Speaker note:</b> " + escape(item["notes"]) + "</p></section>")
        for line in lines:
            pdf_story.append(Paragraph(escape(line), body))
        pdf_story.append(Paragraph("Speaker note: " + escape(item["notes"]), small))
        box(slide, .6, top, 12.1, 6.65 - top, "\n".join(lines), 18 if item.get("rows") else 21)
        box(slide, .6, 6.95, 12.1, .4, f"FISERV / ATL105 / OCTOBER 5 RUN1 | REVIEW REQUIRED | {index + 1} / {len(content)}", 10, False, "626B6C")
        slide.notes_slide.notes_text_frame.text = item["notes"]
    markdown.extend([
        "## BR Meaning and Supporting-Chain Validation", "",
        "We assess these in two separate steps. BR-level equivalence and complete-chain alignment are different conclusions.", "",
        "### 1. Compare BR Meaning", "",
        "Check whether AI BR(s) cover the independent Test BR's applicability and business context, triggering conditions, required behavior and outcome, and exceptions and restrictions.", "",
        "Both interpretations must be supported by the specification. Similar wording or shared IDs is not enough.", "",
        "### 2. Validate the Supporting Chain", "",
        "Then check whether the linked artifacts actually implement that meaning.", "",
        "| Artifact | Question |", "|---|---|",
        "| TS | Does the scenario exercise the BR's condition and behavior? |",
        "| TC | Does the case assert the required result? |",
        "| TD | Does the data activate the intended condition and preserve dependencies? |", "",
        "### Illustrative Lifecycle Example", "",
        "Suppose both BRs say: Completion must reuse the authorization's Sequence Number.", "",
        "AI could use 000001 for both messages; Test could use 100001 for both. Different values are acceptable because reuse within each transaction pair is preserved.", "",
        "If the AI case checks only that Sequence Number has six digits, its chain proves formatting, not lifecycle reuse.", "",
        "This is an illustrative example, not a claim of an approved lifecycle match in the assessed delivery. The actual 000001 versus 100001 comparison below establishes six-digit formatting only.", "",
        "### Interpret the Result", "",
        "- Equivalent BR statements + inadequate tests: BR-level match, incomplete test coverage.",
        "- Complete links + wrong business assertion: structurally complete, not semantically covered.",
        "- Equivalent BRs + aligned TS/TC/TD: aligned full-chain design, still not proof of successful host execution.", "",
    ])
    html.append('''<section id="br-meaning-and-chain"><h2>BR Meaning and Supporting-Chain Validation</h2>
<p>We assess these in <strong>two separate steps</strong>. BR-level equivalence and complete-chain alignment are different conclusions.</p>
<h3>1. Compare BR Meaning</h3><p>Check whether AI BR(s) cover the independent Test BR's:</p>
<ul><li>Applicability and business context.</li><li>Triggering conditions.</li><li>Required behavior and outcome.</li><li>Exceptions and restrictions.</li></ul>
<p>Both interpretations must be supported by the specification. Similar wording or shared IDs is not enough.</p>
<h3>2. Validate the Supporting Chain</h3><p>Then check whether the linked artifacts actually implement that meaning:</p>
<div class="table-wrap"><table><thead><tr><th>Artifact</th><th>Question</th></tr></thead><tbody>
<tr><td>TS</td><td>Does the scenario exercise the BR's condition and behavior?</td></tr>
<tr><td>TC</td><td>Does the case assert the required result?</td></tr>
<tr><td>TD</td><td>Does the data activate the intended condition and preserve dependencies?</td></tr>
</tbody></table></div>
<h3>Illustrative Lifecycle Example</h3><p>Suppose both BRs say: <strong>Completion must reuse the authorization's Sequence Number.</strong></p>
<p>AI could use <code>000001</code> for both messages; Test could use <code>100001</code> for both. Different values are acceptable because the business relationship, <strong>reuse within each transaction pair</strong>, is preserved.</p>
<p>However, if the AI case checks only <strong>Sequence Number has six digits</strong>, its chain proves formatting, <strong>not lifecycle reuse</strong>.</p>
<p class="scope">This is an illustrative example, not a claim of an approved lifecycle match in the assessed delivery. The actual <code>000001</code> versus <code>100001</code> comparison below establishes six-digit formatting only.</p>
<h3>Interpret the Result</h3><ul>
<li>Equivalent BR statements + inadequate tests: <strong>BR-level match, incomplete test coverage.</strong></li>
<li>Complete links + wrong business assertion: <strong>structurally complete, not semantically covered.</strong></li>
<li>Equivalent BRs + aligned TS/TC/TD: <strong>aligned full-chain design</strong>, still not proof of successful host execution.</li>
</ul></section>''')
    markdown.extend(["## Actual Artifact Comparison Appendix", "", "Examples are drawn from frozen AI artifacts and the independent Test Solution populated aggregate. No full-chain perfect match is certified.", ""])
    html.append("<section><h2>Actual Artifact Comparison Appendix</h2>")
    for example in proof["examples"]:
        markdown.extend(["### " + example["title"], "", "| Layer | AI Solution | Test Solution |", "|---|---|---|"])
        comparison = [[key, example["ai"][key], example["test"][key]] for key in ["BR", "TS", "TC", "TD"]]
        comparison.extend([["BR statement", example["ai"]["brStatement"], example["test"]["brStatement"]], ["Field value", example["ai"]["value"], example["test"]["value"]], ["Outcome", str(example["ai"]["expectedResponse"]), "Validator " + str(example["test"]["expectedValidation"])]])
        markdown.extend("| " + " | ".join(str(cell).replace("\n", " ").replace("|", "\\|") for cell in row) + " |" for row in comparison)
        markdown.extend(["", "Disposition: " + example["disposition"] + ". Whole-chain confirmed: false.", ""])
        html.append("<h3>" + escape(example["title"]) + '</h3><div class="table-wrap"><table>' + "".join("<tr>" + "".join("<td>" + escape(str(cell)) + "</td>" for cell in row) + "</tr>" for row in comparison) + "</table></div><p>Disposition: " + escape(example["disposition"]) + "; whole-chain confirmed: false.</p>")
    markdown.extend(["## Evidence and Demo Links", "", "- [Run report and filter views](../AI-ARTIFACT-FILTER-VIEW.html)", "- [Semantic report and scoped cases](../SEMANTIC-FIRST-BATCH-REPORT.html)", "- [Complete independent matrix CSV](../reconstructed-requirement-matrix.csv)", "- [Presentation facts and hashes](presentation-facts.json)", "- [Example artifact extracts and hashes](presentation-examples.json)", "", "Do not present September candidates, placeholder chains, metadata companions or source-rule inventories as approved October full-chain matches."])
    markdown.extend(["", "## Five-Minute Demo Checklist", "", "1. Open the Run report: show 12,679 scenarios, 21,123 cases and the execution boundary.", "2. Open No-TC scenarios: select Transaction unresolved (2,026), then Response deferred (1,706). Reset between demonstrations.", "3. Export selected CSV; explain that these filters cover the 3,732 no-TC backlog only.", "4. Open Coverage: select Segment-wise linkage and point to Segment 100; label it structural linkage, not semantic coverage.", "5. Open the semantic report: select negative intent (25) or source-predicate failure (2); show one case detail and source evidence.", "6. Open Code-1 queue: select Financial Transaction Request (416 candidates); expected TC responses remain empty.", "7. Show Matrix & history: complete reconstruction, 63 unresolved rows and retained previous reports.", "8. Show the comparison appendix: 000001 versus 100001 satisfies the same six-digit predicate, but full-chain equivalence is unconfirmed.", "", "## Statements to Avoid", "", "- '53.13% of the business rules are independently validated.' It is producer-internal linkage.", "- '522 perfect matches.' Those are September atomic candidates, not October full-chain confirmations.", "- 'The Test Solution has 1,326 approved BRs.' The structural package includes review-required catalog records and placeholders.", "- 'Different TD values always match.' Literal constants and context-sensitive values cannot be freely substituted.", "- '40 passed tests prove production readiness.' These are synthetic isolated predicate probes, not executed host transactions."])
    html.append('<p><a href="presentation-facts.json">Facts and source hashes</a> | <a href="presentation-examples.json">Actual example extracts and hashes</a> | <a href="../reconstructed-requirement-matrix.csv">Full matrix CSV</a></p></section></main></html>')
    html[0] = html[0].replace("body{margin:0;", "body{overflow-wrap:anywhere;margin:0;")
    (output / "FISERV-BUSINESS-BRIEF.md").write_text("\n".join(markdown), encoding="utf-8")
    (output / "FISERV-BUSINESS-BRIEF.html").write_text("".join(html), encoding="utf-8")
    if not html_only:
        deck.save(str(output / "FISERV-BUSINESS-BRIEF.pptx"))
        SimpleDocTemplate(str(output / "FISERV-BUSINESS-BRIEF.pdf"), pagesize=landscape(A4), leftMargin=40, rightMargin=40, topMargin=32, bottomMargin=32).build(pdf_story)


def self_test():
    empty = {value: [] for value in ARTIFACTS.values()}
    assert counts(empty) == {"BR": 0, "TS": 0, "TC": 0, "TD": 0}
    invalid = {**empty, "businessRequirements": [{"id": "BR-1"}, {"id": "BR-1"}]}
    try:
        counts(invalid)
    except ValueError:
        pass
    else:
        raise AssertionError("Duplicate canonical artifacts must not inflate presentation totals")
    print("PASS: canonical inventory counting and duplicate/missing-ID refusal")


def validate_outputs(review):
    from html.parser import HTMLParser
    from pptx import Presentation
    from pypdf import PdfReader
    output = native(review / "business-presentation-2026-10-07")
    facts = read(output / "presentation-facts.json")
    proof = read(output / "presentation-examples.json")
    for relative, expected in facts["inputSha256"].items():
        if digest(MODULE / relative) != expected:
            raise ValueError("Presentation input changed: " + relative)
    for file, expected in proof["inputSha256"].items():
        if digest(Path(file)) != expected:
            raise ValueError("Example source changed: " + file)
    if facts["ai"] != {"BR": 6887, "TS": 12679, "TC": 21123, "logicalCasesWithTD": 21096, "physicalTDRequestFiles": 21210, "physicalRequestMetadataPairs": 21210}:
        raise ValueError("AI presentation figures changed")
    if facts["testSolutionPopulatedAggregate"] != {"BR": 725, "TS": 1343, "TC": 1424, "TD": 1416}:
        raise ValueError("Populated Test Solution presentation figures changed")
    if len(proof["examples"]) != 3 or any(item["wholeChainConfirmed"] for item in proof["examples"]):
        raise ValueError("Example scope incorrectly promoted")
    if (proof["examples"][0]["ai"]["value"], proof["examples"][0]["test"]["value"]) != ("000001", "100001"):
        raise ValueError("Different-but-valid sequence example changed")
    deck = Presentation(str(output / "FISERV-BUSINESS-BRIEF.pptx"))
    if len(deck.slides) != 10:
        raise ValueError("Expected ten presentation slides")
    deck_text = []
    for slide in deck.slides:
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError("A presentation slide lacks speaker notes")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width or shape.top + shape.height > deck.slide_height:
                raise ValueError("Presentation shape exceeds slide bounds")
            if shape.has_text_frame:
                deck_text.append(shape.text)
            if shape.has_table:
                deck_text.extend(cell.text for row in shape.table.rows for cell in row.cells)
    combined = "\n".join(deck_text)
    for marker in ["1 |", "2 |", "3 |", "4 |", "5 |", "6A |", "6B |", "7 |", "6,887", "21,210", "725", "1,343", "000001", "100001"]:
        if marker not in combined:
            raise ValueError("Missing presentation question/statistic: " + marker)
    pdf = PdfReader(output / "FISERV-BUSINESS-BRIEF.pdf")
    pdf_text = " ".join(" ".join(page.extract_text() or "" for page in pdf.pages).split())
    for marker in ["6,887", "21,210", "61,044", "000001", "100001", "September", "full-chain"]:
        if marker not in pdf_text:
            raise ValueError("PDF lacks required evidence/scope: " + marker)
    class Links(HTMLParser):
        def __init__(self):
            super().__init__()
            self.targets = []

        def handle_starttag(self, tag, attributes):
            if tag == "a":
                self.targets.extend(value for key, value in attributes if key == "href")
    links = Links()
    brief_html = (output / "FISERV-BUSINESS-BRIEF.html").read_text(encoding="utf-8")
    links.feed(brief_html)
    for marker in ["br-meaning-and-chain", "1. Compare BR Meaning", "2. Validate the Supporting Chain",
                   "reuse within each transaction pair", "not lifecycle reuse", "BR-level match, incomplete test coverage",
                   "structurally complete, not semantically covered", "still not proof of successful host execution"]:
        if marker not in brief_html:
            raise ValueError("Missing BR meaning/chain distinction: " + marker)
    for target in links.targets:
        if not native(output / target).is_file():
            raise ValueError("Broken presentation/report link: " + target)
    prior_file = output / "presentation-verification.json"
    prior = read(prior_file) if prior_file.is_file() else {}
    native_unchanged = prior.get("officeRenderingVerified", False) and prior.get("outputSha256", {}).get("FISERV-BUSINESS-BRIEF.pptx") == digest(output / "FISERV-BUSINESS-BRIEF.pptx")
    verification = {"status": "PASS", "slides": len(deck.slides), "pdfPages": len(pdf.pages), "questionsCovered": 7,
                    "realArtifactExamples": 3, "speakerNotesPresent": True, "sourceHashesVerified": True,
                    "slideShapeBoundsVerified": True, "officeRenderingVerified": native_unchanged,
                    "outputSha256": {file.name: digest(file) for file in output.iterdir() if file.is_file() and file.name != "presentation-verification.json"}}
    if native_unchanged:
        for key in ["nativeRenderedSlides", "nativeTextOverflowFindings", "nativeRenderer"]:
            verification[key] = prior[key]
    (output / "presentation-verification.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
    print("PASS: 10 slides with notes, all 7 questions, 3 actual examples, source hashes, shape bounds, PDF content and report links")
    print("PDF pages: " + str(len(pdf.pages)))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("review_directory", nargs="?", type=Path)
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("--html-only", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
    elif args.validate and args.review_directory:
        validate_outputs(args.review_directory)
    elif args.review_directory:
        facts = gather(args.review_directory)
        proof = examples(args.review_directory)
        output = native(args.review_directory / "business-presentation-2026-10-07")
        output.mkdir(parents=True, exist_ok=True)
        (output / "presentation-facts.json").write_text(json.dumps(facts, indent=2, ensure_ascii=True), encoding="utf-8")
        (output / "presentation-examples.json").write_text(json.dumps(proof, indent=2, ensure_ascii=True), encoding="utf-8")
        generate_outputs(output, facts, proof, args.review_directory, args.html_only)
        print(json.dumps({key: value for key, value in facts.items() if key != "inputSha256"}, indent=2))
    else:
        parser.error("review_directory or --self-test required")


if __name__ == "__main__":
    main()
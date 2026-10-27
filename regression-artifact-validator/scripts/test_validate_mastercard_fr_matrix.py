import csv
import importlib.util
import io
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl


SCRIPT = Path(__file__).with_name("validate-mastercard-fr-matrix.py")
SPEC = importlib.util.spec_from_file_location("mastercard_fr_matrix_validation", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
validation = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(validation)


class MastercardFrMatrixValidationTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.curated = validation.read_json(validation.CURATED)
        cls.workbook = (validation.MODULE / cls.curated["artifact"]["path"]).resolve()
        cls.report, cls.outputs, cls.xlsx = validation.build()
        cls.lines = {line["id"]: line for line in cls.report["lines"]}
        cls.scenarios = {item["id"]: item for item in cls.report["scenarios"]}
        cls.findings = {item["id"]: item for item in cls.report["findings"]}

    def test_every_workbook_line_has_exactly_one_assessment(self):
        summary = self.report["summary"]
        self.assertEqual(
            (summary["businessRequirements"], summary["functionalRequirements"], summary["scenarios"],
             summary["validationLines"], summary["expectedResults"]),
            (1, 8, 37, 65, 37))
        self.assertEqual(len(self.lines), summary["validationLines"])
        self.assertEqual(self.report["workbook"]["sha256"], self.curated["artifact"]["sha256"])
        self.assertEqual(self.report["businessRequirement"]["id"], "BR-01")
        self.assertEqual([item["id"] for item in self.report["scenarios"]], [f"TS-{n:03d}" for n in range(1, 38)])

    def test_element_citations_are_checked_against_chapter_13(self):
        check = {item["line"]: item for item in self.report["elementCheck"]}
        mismatches = {line: (item["citedElement"], item["citedName"], item["correctElement"])
                      for line, item in check.items() if item["result"] == "MISMATCH"}
        self.assertEqual(mismatches, {
            "TS-002.V1": (51, "Mail Text Data", 12), "TS-010.V1": (24, "Data Type Indicator", 2),
            "TS-016.V1": (80, "Purchase Code", 86), "TS-022.V1": (6, "Approved Amount", 5)})
        self.assertEqual({line for line, item in check.items() if item["result"] == "MATCH"}, {"TS-007.V1", "TS-008.V1"})

    def test_contradicted_lines_reject_their_scenarios_and_the_artifact(self):
        contradicted = {line_id for line_id, line in self.lines.items() if line["verdict"] == "CONTRADICTED"}
        self.assertEqual(contradicted, {"TS-002.V1", "TS-004.V1", "TS-010.V1", "TS-016.V1", "TS-022.V1"})
        for line_id in contradicted:
            self.assertEqual(self.lines[line_id]["severity"], "HIGH")
            self.assertTrue(self.lines[line_id]["correction"])
            self.assertEqual(self.scenarios[line_id.split(".")[0]]["outcome"], "REJECTED_NOT_INTAKE_READY")
        self.assertEqual(self.report["outcome"], "REJECTED_NOT_INTAKE_READY")
        self.assertFalse(self.report["executionCertified"])
        self.assertEqual(self.report["status"], "DRAFT_REVIEW_REQUIRED")
        self.assertNotIn("ACCEPTED", {item["outcome"] for item in self.report["scenarios"]},
                         "no scenario can be ACCEPTED while every expected result needs execution evidence")

    def test_evidence_resolves_and_review_lines_carry_open_questions(self):
        for line in self.lines.values():
            self.assertTrue(line["evidence"] or line["testSolutionBrs"] or line["draftBrs"], line["id"])
            for item in line["evidence"]:
                if "absence" not in item:
                    self.assertGreater(item["line"], 0)
                    self.assertEqual(len(item["sha256"]), 64)
            if line["verdict"] in validation.REVIEW_VERDICTS:
                self.assertIn("REVIEW_REQUIRED", line["classes"], line["id"])
                self.assertTrue(line["questions"], line["id"])
        self.assertEqual(self.lines["TS-011.V3"]["evidence"][0]["absence"]["occurrences"], 0)

    def test_test_solution_status_is_read_live_and_gaps_are_computed(self):
        for line in self.lines.values():
            states = {item["state"] for item in line["testSolutionBrs"]}
            self.assertEqual(line["coverage"], validation.line_coverage(states, [d["id"] for d in line["draftBrs"]]))
            if line["verdict"] in validation.SUPPORTED_VERDICTS and line["coverage"] in ("NONE", "DRAFT_ONLY", "OUTSIDE_MC_CONTEXT"):
                self.assertIn("TEST_SOLUTION_GAP", line["classes"], line["id"])
        self.assertEqual(self.lines["TS-009.V1"]["coverage"], "TRAINED_CHAIN")
        self.assertIn("TEST_SOLUTION_GAP", self.lines["TS-028.V1"]["classes"])
        mc = validation.load_module("mc_matrix_for_test", "build-mastercard-fr-matrix.py")
        t03, t04 = self.findings["MCV-T03"], self.findings["MCV-T04"]
        self.assertEqual(t03["status"] == "OPEN", bool(self.report["multiSegmentScopeExclusions"]))
        for item in self.report["multiSegmentScopeExclusions"]:
            self.assertTrue(any("," in segment for segment in item["segments"]), item["id"])
        self.assertEqual(t04["status"] == "OPEN", "BR-APPG-CODE-3" in mc.TEMPLATE_MAP["TS-020"][0])

    def test_transcription_differences_are_verified_against_both_sources(self):
        differences = {item["ts"]: item for item in self.report["transcriptionDifferences"]}
        self.assertEqual(set(differences), {"TS-001", "TS-034", "TS-035", "TS-036"})
        self.assertTrue(all(item["verified"] for item in differences.values()))
        self.assertEqual(self.report["baFidelityChecks"], [])

    def test_changed_workbook_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            changed = Path(folder) / "matrix.xlsx"
            shutil.copyfile(validation.long_path(self.workbook), changed)
            book = openpyxl.load_workbook(changed)
            sheet = book[self.curated["artifact"]["sheet"]]
            cell = next(c for row in sheet.iter_rows() for c in row if c.value == "Element 80 present")
            cell.value = "Element 86 present"
            book.save(changed)
            with self.assertRaisesRegex(ValueError, "SHA-256"):
                validation.build_report(changed)
            with patch.object(validation, "sha256", return_value=self.curated["artifact"]["sha256"]):
                with self.assertRaisesRegex(ValueError, "differ from the curated assessment"):
                    validation.build_report(changed)

    def test_inconsistent_curated_assessments_are_refused(self):
        line = dict(next(item for item in self.curated["lines"] if item["verdict"] == "CONTRADICTED"))
        line["severity"] = "LOW"
        with self.assertRaisesRegex(ValueError, "must be HIGH"):
            validation.check_curated_line(line)
        external = dict(next(item for item in self.curated["lines"] if item["verdict"] == "EXTERNAL_POLICY"))
        external["questions"] = []
        with self.assertRaisesRegex(ValueError, "open question"):
            validation.check_curated_line(external)

    def test_outputs_are_complete_current_and_injection_safe(self):
        self.assertEqual(tuple(self.outputs), validation.TEXT_OUTPUTS)
        rows = list(csv.DictReader(io.StringIO(self.outputs["line-validation.csv"])))
        self.assertEqual(len(rows), 65)
        for row in rows:
            for value in row.values():
                self.assertFalse(value.startswith(("=", "+", "-", "@")), value)
        book = openpyxl.load_workbook(io.BytesIO(self.xlsx))
        self.assertEqual(book.sheetnames, ["Summary", "Line Validation", "Scenarios", "Findings", "Element Check",
                                           "BA Transcription"])
        self.assertEqual(book["Line Validation"].max_row, 66)
        for name, text in self.outputs.items():
            self.assertEqual(validation.read_bytes(validation.OUT_DIR / name).decode("utf-8"), text,
                             f"{name} is stale; run scripts/validate-mastercard-fr-matrix.py")


if __name__ == "__main__":
    unittest.main()
